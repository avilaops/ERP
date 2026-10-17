import io
import unittest
import zipfile
import shutil
from pathlib import Path
from extract import extract


class ExtractionTests(unittest.TestCase):
    def test_html_does_not_render_scripts_and_retains_links(self):
        item = extract(b'<html><title>Manual</title><main><p>Documento oficial</p><a href="a.pdf">PDF</a><script>steal()</script></main></html>', 'a.html', 'text/html')
        self.assertNotIn('steal', item['pages'][0]['content'])
        self.assertEqual(item['links'], ['a.pdf'])

    def test_xsd_catalog_has_types_enums_and_original(self):
        xml = b'<xs:schema xmlns:xs="http://www.w3.org/2001/XMLSchema" targetNamespace="urn:nfe"><xs:element name="cStat" type="xs:string" minOccurs="0"/><xs:simpleType name="Status"><xs:restriction base="xs:string"><xs:enumeration value="100"/></xs:restriction></xs:simpleType><xs:include schemaLocation="tipos.xsd"/></xs:schema>'
        item = extract(xml, 'nfe.xsd', 'application/xml')
        self.assertEqual(item['elements'][0]['namespace'], 'urn:nfe')
        self.assertEqual(item['elements'][0]['min_occurs'], '0')
        self.assertEqual(item['elements'][1]['details']['enumerations'], ['100'])
        self.assertIn('tipos.xsd', item['links'])

    def test_external_xml_entity_is_rejected(self):
        with self.assertRaises(Exception):
            extract(b'<!DOCTYPE x [<!ENTITY secret SYSTEM "file:///etc/passwd">]><x>&secret;</x>', 'a.xml', 'application/xml')

    def test_zip_paths_and_expansion_are_rejected(self):
        for name in ('../secret.txt', '/etc/passwd', 'C:/secret.txt'):
            data = io.BytesIO()
            with zipfile.ZipFile(data, 'w') as archive:
                archive.writestr(name, b'hostile')
            with self.assertRaises(ValueError):
                extract(data.getvalue(), 'a.zip')

    def test_sql_is_documentation_only(self):
        item = extract(b'DROP DATABASE erp;', 'unsafe.sql', 'text/plain')
        self.assertFalse(item['metadata']['sql_executed'])
        self.assertEqual(item['metadata']['sql_statement_types'], ['DROP'])

    def test_nested_zip_is_catalogued_with_bounded_expansion(self):
        inner = io.BytesIO()
        with zipfile.ZipFile(inner, 'w') as archive:
            archive.writestr('schema.xsd', '<xs:schema xmlns:xs="http://www.w3.org/2001/XMLSchema"><xs:element name="cStat"/></xs:schema>')
        outer = io.BytesIO()
        with zipfile.ZipFile(outer, 'w') as archive:
            archive.writestr('package.zip', inner.getvalue())
        result = extract(outer.getvalue(), 'outer.zip')
        self.assertEqual(result['elements'][0]['name'], 'cStat')
        self.assertEqual(result['warnings'], [])

    def test_zip_schema_include_is_not_mistaken_for_remote_download(self):
        data=io.BytesIO()
        with zipfile.ZipFile(data,'w') as archive:
            archive.writestr('schema.xsd','<xs:schema xmlns:xs="http://www.w3.org/2001/XMLSchema"><xs:include schemaLocation="tipos.xsd"/></xs:schema>')
            archive.writestr('tipos.xsd','<xs:schema xmlns:xs="http://www.w3.org/2001/XMLSchema"/>')
        item=extract(data.getvalue(),'schemas.zip')
        self.assertEqual(item['links'],[])
        self.assertTrue(item['metadata']['schema_dependencies'][0]['in_archive'])

    def test_spreadsheet_formatting_does_not_expand_data_dimensions(self):
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill
        book = Workbook()
        sheet = book.active
        sheet.append(['CFOP', 'Descrição'])
        sheet.append(['5102', 'Venda'])
        sheet.cell(2, 1000).fill = PatternFill(fill_type='solid', fgColor='FFFF00')
        data = io.BytesIO()
        book.save(data)
        item = extract(data.getvalue(), 'cfop.xlsx')
        self.assertIn('5102\tVenda', item['pages'][0]['content'])

    @unittest.skipUnless(shutil.which('tesseract') or Path('C:/Program Files/Tesseract-OCR/tesseract.exe').exists(), 'Tesseract não instalado')
    def test_scanned_pdf_uses_real_portuguese_ocr(self):
        from PIL import Image, ImageDraw, ImageFont
        image = Image.new('RGB', (1600, 600), 'white')
        draw = ImageDraw.Draw(image)
        draw.text((80, 150), 'DOCUMENTO FISCAL TESTE', font=ImageFont.load_default(size=70), fill='black')
        pdf = io.BytesIO()
        image.save(pdf, 'PDF', resolution=150)
        result = extract(pdf.getvalue(), 'scanned.pdf', 'application/pdf')
        self.assertIn('FISCAL', result['pages'][0]['content'].upper())
        self.assertEqual(result['metadata']['ocr'][0]['method'], 'tesseract-por-tsv')


if __name__ == '__main__':
    unittest.main()
